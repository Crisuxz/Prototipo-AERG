"""Validacion y extraccion de texto de los trabajos entregados (RF-03, RF-04, RF-05).

Solo se aceptan TXT, PDF y DOCX. El nombre original nunca se usa como ruta: se saneia
para trazabilidad y el archivo se guarda con un nombre generado (uuid4 + extension),
dentro de `UPLOAD_DIR` (Parte 66)."""

import io
import re
import unicodedata
import uuid
from dataclasses import dataclass
from pathlib import Path

import docx
from pypdf import PdfReader

from app.core.errors import DomainError

ALLOWED_EXTENSIONS = {"txt", "pdf", "docx"}
ALLOWED_MIME_TYPES = {
    "txt": {"text/plain", "application/octet-stream", ""},
    "pdf": {"application/pdf", "application/octet-stream", ""},
    "docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
        "",
    },
}
FILE_SIGNATURES = {"pdf": b"%PDF-", "docx": b"PK\x03\x04"}


@dataclass
class ExtractedDocument:
    text: str
    extension: str
    stored_filename: str
    safe_original_filename: str
    size_bytes: int


def sanitize_filename(filename: str) -> str:
    """Neutraliza path traversal y caracteres raros; solo para mostrar/trazar, nunca como ruta."""
    base = Path(filename.replace("\\", "/")).name
    normalized = unicodedata.normalize("NFKD", base).encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r"[^A-Za-z0-9._-]", "_", normalized).lstrip(".")
    return cleaned[:255] or "archivo"


def get_extension(filename: str) -> str:
    return Path(sanitize_filename(filename)).suffix.lower().lstrip(".")


def normalize_text(raw: str) -> str:
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_txt(content: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DomainError("EXTRACTION_FAILED", "No fue posible decodificar el archivo de texto.", 500)


def extract_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise DomainError(
            "EXTRACTION_FAILED", "El PDF esta danado o no se pudo leer su contenido.", 500
        ) from exc


def extract_docx(content: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    except Exception as exc:
        raise DomainError(
            "EXTRACTION_FAILED", "El DOCX esta danado o no se pudo leer su contenido.", 500
        ) from exc


EXTRACTORS = {"txt": extract_txt, "pdf": extract_pdf, "docx": extract_docx}


class DocumentProcessingService:
    def __init__(self, max_upload_mb: int, upload_dir: str):
        self.max_upload_bytes = max_upload_mb * 1024 * 1024
        self.max_upload_mb = max_upload_mb
        self.upload_dir = Path(upload_dir)

    def process(self, filename: str, content: bytes, mime_type: str | None) -> ExtractedDocument:
        extension = get_extension(filename)
        if extension not in ALLOWED_EXTENSIONS:
            raise DomainError(
                "INVALID_FILE_TYPE",
                "Formato no soportado. Solo se aceptan archivos TXT, PDF o DOCX.",
                400,
            )

        declared_mime = (mime_type or "").split(";")[0].strip().lower()
        if declared_mime and declared_mime not in ALLOWED_MIME_TYPES[extension]:
            raise DomainError(
                "INVALID_FILE_TYPE",
                "El tipo de contenido del archivo no corresponde con su extension.",
                400,
            )

        if len(content) > self.max_upload_bytes:
            raise DomainError(
                "FILE_TOO_LARGE",
                "El archivo supera el limite de " + str(self.max_upload_mb) + " MB.",
                413,
            )

        signature = FILE_SIGNATURES.get(extension)
        if signature and not content.startswith(signature):
            raise DomainError(
                "INVALID_FILE_TYPE",
                "El contenido del archivo no corresponde con la extension declarada.",
                400,
            )

        if not content:
            raise DomainError("EMPTY_DOCUMENT", "El archivo esta vacio.", 422)

        text = normalize_text(EXTRACTORS[extension](content))
        if not text:
            raise DomainError(
                "EMPTY_DOCUMENT", "No se encontro texto legible en el documento.", 422
            )

        stored_filename = str(uuid.uuid4()) + "." + extension
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        (self.upload_dir / stored_filename).write_bytes(content)

        return ExtractedDocument(
            text=text,
            extension=extension,
            stored_filename=stored_filename,
            safe_original_filename=sanitize_filename(filename),
            size_bytes=len(content),
        )
